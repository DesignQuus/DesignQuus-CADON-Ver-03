#!/usr/bin/env python3
"""
CADON-BOM SERVER POC - PROMPT 15 & STANDARD CORPORATE EXCEL QUOTATION
Excel Export Engine / Standard Quotation Template Generation & Validation
"""
import sys
import os
import json
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def number_to_korean_won(num) -> str:
    if not num or num <= 0:
        return "금 영 원정"
    units = ["", "만", "억", "조"]
    digits = ["", "일", "이", "삼", "사", "오", "육", "칠", "팔", "구"]
    sub_units = ["", "십", "백", "천"]
    
    num_int = int(round(float(num)))
    str_num = str(num_int)
    result = ""
    unit_idx = 0
    
    while str_num:
        chunk = str_num[-4:]
        str_num = str_num[:-4]
        
        chunk_str = ""
        for i, ch in enumerate(reversed(chunk)):
            d = int(ch)
            if d > 0:
                chunk_str = digits[d] + sub_units[i] + chunk_str
                
        if chunk_str:
            result = chunk_str + units[unit_idx] + " " + result
        unit_idx += 1
        
    return "금 " + result.strip() + " 원정"

def populate_standard_quotation(ws, quote_data: dict):
    font_family = "맑은 고딕"
    
    c_title = "0F172A"      # Slate 900
    c_subtitle = "64748B"   # Slate 500
    c_text_main = "1E293B"  # Slate 800
    c_text_sub = "475569"   # Slate 600
    c_blue = "1E3A8A"       # Deep Blue
    
    fill_header = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")     # Slate 200
    fill_sub_hdr = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")    # Slate 100
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")      # Slate 50
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_line = Side(style='thin', color='CBD5E1')
    thick_line = Side(style='medium', color='334155')
    double_line = Side(style='double', color='334155')
    
    border_all_thin = Border(left=thin_line, right=thin_line, top=thin_line, bottom=thin_line)
    
    def apply_range_border(min_c, min_r, max_c, max_r, border):
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                ws.cell(row=r, column=c).border = border

    # 1. Column Dimensions (B to K: 10 columns)
    ws.column_dimensions['A'].width = 3     # Left margin
    ws.column_dimensions['B'].width = 6     # No
    ws.column_dimensions['C'].width = 16    # 도면번호 (DWG NO)
    ws.column_dimensions['D'].width = 26    # 품명
    ws.column_dimensions['E'].width = 20    # 규격 및 사양
    ws.column_dimensions['F'].width = 13    # 재질
    ws.column_dimensions['G'].width = 9     # 수량
    ws.column_dimensions['H'].width = 8     # 단위
    ws.column_dimensions['I'].width = 16    # 단가
    ws.column_dimensions['J'].width = 18    # 공급가액
    ws.column_dimensions['K'].width = 16    # 비고

    # 2. Document Title
    ws.row_dimensions[2].height = 34
    ws.merge_cells("B2:K2")
    ws["B2"] = "견   적   서"
    ws["B2"].font = Font(name=font_family, size=22, bold=True, color=c_title)
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.row_dimensions[3].height = 16
    ws.merge_cells("B3:K3")
    ws["B3"] = "QUOTATION SPECIFICATION"
    ws["B3"].font = Font(name=font_family, size=9.5, bold=False, color=c_subtitle)
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.row_dimensions[4].height = 8

    # Extract info
    quote = quote_data.get("quote", {})
    customer_name = quote_data.get("customer_name", "고객사 귀하")
    project_name = quote_data.get("project_name", "표준 설비 제작 프로젝트")
    quote_no = quote.get("quote_no", "Q-20260901-002-V1")
    quote_date = quote.get("quote_date", "2026-09-04")
    if "T" in str(quote_date):
        quote_date = str(quote_date).split("T")[0]
        
    supplier = quote_data.get("supplier", {
        "business_no": "",
        "company_name": "",
        "ceo_name": "",
        "address": "",
        "biz_type": "",
        "biz_category": "",
        "tel": "",
        "fax": "",
        "email": "",
        "manager": ""
    })
    
    terms = quote_data.get("terms", {
        "delivery_terms": "발주 확정 후 30일 이내 납품 (도면 승인 기준)",
        "payment_terms": "세금계산서 발행 후 30일 이내 현금 결제 (협의 가능)",
        "validity_terms": "견적 제출일로부터 30일간 유효",
        "bank_account": "",
        "remarks": "1. 본 견적서는 CAD 도면 정밀 분석 기반 표준 산출 견적서입니다.\n2. 사양 변경 시 견적 금액이 변동될 수 있습니다."
    })

    # 3. Dual Header: Left (Client) / Right (Supplier)
    for r in range(5, 12):
        ws.row_dimensions[r].height = 20

    # Left Box: 공급받는 자 (Cols B to E, Rows 5 to 11)
    ws.merge_cells("B5:E5")
    ws["B5"] = "공 급 받 는  자  (발 주 처)"
    ws["B5"].font = Font(name=font_family, size=10, bold=True, color=c_title)
    ws["B5"].fill = fill_sub_hdr
    ws["B5"].alignment = Alignment(horizontal="center", vertical="center")
    
    cust_display = str(customer_name)
    if not cust_display.endswith("귀하"):
        cust_display += " 귀하"

    client_fields = [
        ("발 주 처", cust_display, True),
        ("프 로 젝트", project_name, False),
        ("견 적 일 자", quote_date, False),
        ("견 적 번 호", quote_no, True),
        ("납 기 일 자", terms.get("delivery_terms", "발주 후 30일 이내"), False),
        ("유 효 기 간", terms.get("validity_terms", "제출일로부터 30일간"), False),
    ]
    
    for idx, (lbl, val, is_bold) in enumerate(client_fields, start=6):
        ws.cell(row=idx, column=2, value=lbl).font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
        ws.cell(row=idx, column=2).fill = fill_zebra
        ws.cell(row=idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
        
        ws.merge_cells(start_row=idx, start_column=3, end_row=idx, end_column=5)
        cell_val = ws.cell(row=idx, column=3, value=val)
        cell_val.font = Font(name=font_family, size=9.5 if not is_bold else 10, bold=is_bold, color=c_title)
        cell_val.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    apply_range_border(2, 5, 5, 11, border_all_thin)

    # Right Box: 공급자 (Cols F to K, Rows 5 to 11)
    ws.merge_cells("F5:F11")
    ws["F5"] = "공\n\n급\n\n자"
    ws["F5"].font = Font(name=font_family, size=10, bold=True, color=c_title)
    ws["F5"].fill = fill_sub_hdr
    ws["F5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Supplier rows
    ws["G5"] = "등록번호"
    ws["G5"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G5"].fill = fill_zebra
    ws["G5"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("H5:K5")
    ws["H5"] = supplier.get("business_no", "")
    ws["H5"].font = Font(name=font_family, size=10, bold=True, color=c_title)
    ws["H5"].alignment = Alignment(horizontal="center", vertical="center")

    ws["G6"] = "상    호"
    ws["G6"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G6"].fill = fill_zebra
    ws["G6"].alignment = Alignment(horizontal="center", vertical="center")
    ws["H6"] = supplier.get("company_name", "")
    ws["H6"].font = Font(name=font_family, size=9.5, bold=True, color=c_title)
    ws["H6"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["I6"] = "대 표 자"
    ws["I6"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["I6"].fill = fill_zebra
    ws["I6"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("J6:K6")
    ceo_text = supplier.get('ceo_name', '')
    ws["J6"] = f"{ceo_text}  (인)" if ceo_text else "(인)"
    ws["J6"].font = Font(name=font_family, size=9.5, bold=True, color=c_title)
    ws["J6"].alignment = Alignment(horizontal="center", vertical="center")

    ws["G7"] = "소 재 지"
    ws["G7"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G7"].fill = fill_zebra
    ws["G7"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("H7:K7")
    ws["H7"] = supplier.get("address", "")
    ws["H7"].font = Font(name=font_family, size=8.5, color=c_text_main)
    ws["H7"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["G8"] = "업    태"
    ws["G8"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G8"].fill = fill_zebra
    ws["G8"].alignment = Alignment(horizontal="center", vertical="center")
    ws["H8"] = supplier.get("biz_type", "")
    ws["H8"].font = Font(name=font_family, size=9, color=c_text_main)
    ws["H8"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["I8"] = "종    목"
    ws["I8"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["I8"].fill = fill_zebra
    ws["I8"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("J8:K8")
    ws["J8"] = supplier.get("biz_category", "")
    ws["J8"].font = Font(name=font_family, size=9, color=c_text_main)
    ws["J8"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["G9"] = "전화번호"
    ws["G9"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G9"].fill = fill_zebra
    ws["G9"].alignment = Alignment(horizontal="center", vertical="center")
    ws["H9"] = supplier.get("tel", "")
    ws["H9"].font = Font(name=font_family, size=9, color=c_text_main)
    ws["H9"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["I9"] = "팩스번호"
    ws["I9"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["I9"].fill = fill_zebra
    ws["I9"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("J9:K9")
    ws["J9"] = supplier.get("fax", "")
    ws["J9"].font = Font(name=font_family, size=9, color=c_text_main)
    ws["J9"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["G10"] = "담 당 자"
    ws["G10"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G10"].fill = fill_zebra
    ws["G10"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("H10:K10")
    ws["H10"] = supplier.get("manager", "")
    ws["H10"].font = Font(name=font_family, size=9, color=c_text_main)
    ws["H10"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws["G11"] = "이 메 일"
    ws["G11"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws["G11"].fill = fill_zebra
    ws["G11"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("H11:K11")
    ws["H11"] = supplier.get("email", "")
    ws["H11"].font = Font(name=font_family, size=9, color=c_text_main)
    ws["H11"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    apply_range_border(6, 5, 11, 11, border_all_thin)

    # 4. Grand Total Banner (Row 13) - Printer-Friendly clean border and off-white background
    ws.row_dimensions[12].height = 8
    ws.row_dimensions[13].height = 28
    
    total_amount = float(quote.get("total_amount", 0.0))
    korean_won = number_to_korean_won(total_amount)
    
    ws.merge_cells("B13:G13")
    ws["B13"] = f"견적 총액 (VAT 포함):   {korean_won}"
    ws["B13"].font = Font(name=font_family, size=11, bold=True, color=c_title)
    ws["B13"].fill = fill_zebra
    ws["B13"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("H13:K13")
    ws["H13"] = f"₩ {int(total_amount):,}"
    ws["H13"].font = Font(name=font_family, size=14, bold=True, color=c_blue)
    ws["H13"].fill = fill_zebra
    ws["H13"].alignment = Alignment(horizontal="right", vertical="center", indent=1)

    for col in range(2, 12):
        cell = ws.cell(row=13, column=col)
        cell.border = Border(top=thick_line, bottom=thick_line, left=thin_line, right=thin_line)
        cell.fill = fill_zebra

    # Row 14: Notice
    ws.row_dimensions[14].height = 18
    ws.merge_cells("B14:K14")
    ws["B14"] = "※ 아래와 같이 견적서를 제출하오니 검토 후 재가하여 주시기 바랍니다."
    ws["B14"].font = Font(name=font_family, size=9, italic=True, color=c_subtitle)
    ws["B14"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # 5. Table Header (Row 15)
    ws.row_dimensions[15].height = 24
    table_headers = [
        ("No", 2),
        ("도면번호", 3),
        ("품명 (Item Description)", 4),
        ("규격 및 사양 (Specification)", 5),
        ("재질", 6),
        ("수량", 7),
        ("단위", 8),
        ("단가 (원)", 9),
        ("공급가액 (원)", 10),
        ("비고", 11),
    ]
    for h_text, col_idx in table_headers:
        cell = ws.cell(row=15, column=col_idx, value=h_text)
        cell.font = Font(name=font_family, size=9.5, bold=True, color=c_title)
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thick_line, bottom=thick_line, left=thin_line, right=thin_line)

    # 6. Item Rows (Row 16 to 15 + N)
    items = quote_data.get("items", [])
    start_row = 16
    
    for idx, it in enumerate(items):
        row = start_row + idx
        ws.row_dimensions[row].height = 20
        row_fill = fill_zebra if (idx % 2 == 1) else fill_white
        
        c_no = ws.cell(row=row, column=2, value=idx + 1)
        c_no.alignment = Alignment(horizontal="center", vertical="center")
        
        dwg_no = it.get("drawing_no") or it.get("master_code") or "-"
        c_dwg = ws.cell(row=row, column=3, value=dwg_no)
        c_dwg.alignment = Alignment(horizontal="center", vertical="center")
        
        c_name = ws.cell(row=row, column=4, value=it.get("item_name", ""))
        c_name.font = Font(name=font_family, size=9.5, bold=True, color=c_title)
        c_name.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        
        c_spec = ws.cell(row=row, column=5, value=it.get("specification", "-"))
        c_spec.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        
        c_mat = ws.cell(row=row, column=6, value=it.get("material", "-"))
        c_mat.alignment = Alignment(horizontal="center", vertical="center")
        
        qty = float(it.get("quantity", 1.0))
        c_qty = ws.cell(row=row, column=7, value=qty)
        c_qty.number_format = "#,##0"
        c_qty.alignment = Alignment(horizontal="right", vertical="center")
        
        c_unit = ws.cell(row=row, column=8, value=it.get("unit", "EA"))
        c_unit.alignment = Alignment(horizontal="center", vertical="center")
        
        uprice = float(it.get("unit_price", 0.0))
        c_price = ws.cell(row=row, column=9, value=uprice)
        c_price.number_format = "#,##0"
        c_price.alignment = Alignment(horizontal="right", vertical="center")
        
        # Live Excel Formula: =G{row}*I{row}
        c_amt = ws.cell(row=row, column=10, value=f"=G{row}*I{row}")
        c_amt.number_format = "#,##0"
        c_amt.font = Font(name=font_family, size=9.5, bold=True, color=c_title)
        c_amt.alignment = Alignment(horizontal="right", vertical="center")
        
        c_rem = ws.cell(row=row, column=11, value=it.get("remark", ""))
        c_rem.alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(2, 12):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = border_all_thin
            cell.fill = row_fill
            if col_idx not in (4, 10):
                cell.font = Font(name=font_family, size=9, color=c_text_main)

    last_item_row = start_row + len(items) - 1 if len(items) > 0 else start_row

    # 7. Summary Breakdown Rows
    tot_row = last_item_row + 1
    ws.row_dimensions[tot_row].height = 22
    
    ws.merge_cells(f"B{tot_row}:F{tot_row}")
    ws[f"B{tot_row}"] = "품목 수량 합계 / 공급가액 소계"
    ws[f"B{tot_row}"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws[f"B{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    # Qty Sum Formula
    ws[f"G{tot_row}"] = f"=SUM(G{start_row}:G{last_item_row})"
    ws[f"G{tot_row}"].font = Font(name=font_family, size=10, bold=True, color=c_blue)
    ws[f"G{tot_row}"].number_format = "#,##0"
    ws[f"G{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    ws[f"H{tot_row}"] = "EA"
    ws[f"H{tot_row}"].font = Font(name=font_family, size=9, bold=True, color=c_text_sub)
    ws[f"H{tot_row}"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws[f"I{tot_row}"] = "공급가액:"
    ws[f"I{tot_row}"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws[f"I{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    # Subtotal Sum Formula
    ws[f"J{tot_row}"] = f"=SUM(J{start_row}:J{last_item_row})"
    ws[f"J{tot_row}"].font = Font(name=font_family, size=10, bold=True, color=c_title)
    ws[f"J{tot_row}"].number_format = "#,##0"
    ws[f"J{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    for c in range(2, 12):
        ws.cell(row=tot_row, column=c).border = Border(top=thick_line, bottom=thin_line, left=thin_line, right=thin_line)
        ws.cell(row=tot_row, column=c).fill = fill_sub_hdr

    # VAT Row
    vat_row = tot_row + 1
    ws.row_dimensions[vat_row].height = 20
    ws.merge_cells(f"B{vat_row}:H{vat_row}")
    
    ws[f"I{vat_row}"] = "부가가치세 (10%):"
    ws[f"I{vat_row}"].font = Font(name=font_family, size=9.5, bold=True, color=c_text_sub)
    ws[f"I{vat_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    ws[f"J{vat_row}"] = f"=ROUND(J{tot_row}*0.1, 0)"
    ws[f"J{vat_row}"].font = Font(name=font_family, size=10, bold=True, color=c_title)
    ws[f"J{vat_row}"].number_format = "#,##0"
    ws[f"J{vat_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    for c in range(2, 12):
        ws.cell(row=vat_row, column=c).border = border_all_thin

    # Grand Total Row
    grand_row = vat_row + 1
    ws.row_dimensions[grand_row].height = 24
    
    ws.merge_cells(f"B{grand_row}:H{grand_row}")
    ws[f"B{grand_row}"] = "총 견적합계 금액 (GRAND TOTAL):"
    ws[f"B{grand_row}"].font = Font(name=font_family, size=10.5, bold=True, color=c_title)
    ws[f"B{grand_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    ws.merge_cells(f"I{grand_row}:K{grand_row}")
    ws[f"I{grand_row}"] = f"=J{tot_row}+J{vat_row}"
    ws[f"I{grand_row}"].font = Font(name=font_family, size=12, bold=True, color=c_blue)
    ws[f"I{grand_row}"].number_format = '"₩"#,##0'
    ws[f"I{grand_row}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    
    for c in range(2, 12):
        ws.cell(row=grand_row, column=c).border = Border(top=thin_line, bottom=double_line, left=thin_line, right=thin_line)
        ws.cell(row=grand_row, column=c).fill = fill_sub_hdr

    # 8. Terms & Notes Box
    terms_start = grand_row + 2
    ws.row_dimensions[terms_start - 1].height = 10
    
    ws.row_dimensions[terms_start].height = 22
    ws.merge_cells(f"B{terms_start}:K{terms_start}")
    ws[f"B{terms_start}"] = "📌 [ 특기사항 및 견적 조건 ]"
    ws[f"B{terms_start}"].font = Font(name=font_family, size=9.5, bold=True, color=c_title)
    ws[f"B{terms_start}"].fill = fill_sub_hdr
    ws[f"B{terms_start}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    term_lines = [
        f"1. 납품 기한: {terms.get('delivery_terms', '발주 후 30일 이내')}",
        f"2. 대금 결제: {terms.get('payment_terms', '세금계산서 발행 후 30일 이내 현금 결제')}",
        f"3. 견적 유효: {terms.get('validity_terms', '견적 제출일로부터 30일간 유효')}",
        f"4. 입금 계좌: {terms.get('bank_account', '기업은행 123-456789-01-012 (예금주: 주식회사 캐드온)')}",
        "5. 기    타: 전자세금계산서는 국세청 홈택스를 통해 자동 발행되며, 도면 사양 변경 시 견적 금액이 변동될 수 있습니다."
    ]
    
    for t_idx, line in enumerate(term_lines, start=terms_start + 1):
        ws.row_dimensions[t_idx].height = 19
        ws.merge_cells(f"B{t_idx}:K{t_idx}")
        ws[f"B{t_idx}"] = line
        ws[f"B{t_idx}"].font = Font(name=font_family, size=8.5, color=c_text_main)
        ws[f"B{t_idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col_idx in range(2, 12):
            ws.cell(row=t_idx, column=col_idx).border = border_all_thin

    apply_range_border(2, terms_start, 11, terms_start + len(term_lines), border_all_thin)

    # 9. Page Setup for Perfect A4 Print
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.views.sheetView[0].showGridLines = True

def create_default_template(filepath: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "견적서"
    sample_data = {
        "quote": {
            "quote_no": "Q-20260901-001",
            "quote_date": "2026-09-01",
            "subtotal": 0,
            "tax_amount": 0,
            "total_amount": 0
        },
        "customer_name": "A기계 (주)",
        "project_name": "2026 자동화 라인 증설",
        "items": []
    }
    populate_standard_quotation(ws, sample_data)
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)

def export_quote_to_excel(quote_data: dict, template_path: str, output_path: str) -> dict:
    start_time = time.time()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "견적서"
    
    populate_standard_quotation(ws, quote_data)
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    
    items = quote_data.get("items", [])
    quote_info = quote_data.get("quote", {})
    
    calc_subtotal = sum(float(it.get("quantity", 1.0)) * float(it.get("unit_price", 0.0)) for it in items)
    calc_vat = round(calc_subtotal * 0.1)
    calc_total = calc_subtotal + calc_vat
    
    server_tot = float(quote_info.get("total_amount", 0.0))
    match = abs(calc_total - server_tot) < 1.0
    
    return {
        "status": "COMPLETED" if match else "VALIDATION_FAILED",
        "output_path": output_path,
        "file_size": os.path.getsize(output_path),
        "exported_items_count": len(items),
        "server_total": server_tot,
        "excel_total": calc_total,
        "is_total_matched": match,
        "duration_ms": int((time.time() - start_time) * 1000)
    }

if __name__ == "__main__":
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    if len(sys.argv) < 4:
        print(json.dumps({"error": "Usage: excel_exporter.py <quote_json> <template_path> <output_path>"}))
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        data = json.load(f)
    res = export_quote_to_excel(data, sys.argv[2], sys.argv[3])
    print(json.dumps(res, ensure_ascii=False, indent=2))

